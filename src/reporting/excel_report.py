from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.analytics.data_loader import AnalyticsDataLoader


class NexusExcelReport:
    """
    Nexus360 Advanced Excel Business Intelligence Engine.

    Architecture
    ------------
    PostgreSQL Warehouse
        -> Analytics Semantic Layer
        -> Python Analytics / ML
        -> Excel Executive BI Workbook
    """

    OUTPUT_DIR = Path("reports/excel")
    PREDICTION_DIR = Path("reports/predictions")

    # ============================================================
    # COLORS
    # ============================================================

    NAVY = "0F172A"
    BLUE = "2563EB"
    LIGHT_BLUE = "DBEAFE"

    GREEN = "16A34A"
    LIGHT_GREEN = "DCFCE7"

    RED = "DC2626"
    LIGHT_RED = "FEE2E2"

    AMBER = "D97706"
    LIGHT_AMBER = "FEF3C7"

    WHITE = "FFFFFF"
    LIGHT_GRAY = "F1F5F9"
    MID_GRAY = "64748B"
    DARK_TEXT = "111827"

    # ============================================================
    # NUMBER FORMATS
    # ============================================================

    CURRENCY_FORMAT = '$#,##0.00;[Red]-$#,##0.00'
    INTEGER_FORMAT = '#,##0'
    DECIMAL_FORMAT = '#,##0.00'
    PERCENT_FORMAT = '0.00%'
    DATE_FORMAT = 'yyyy-mm-dd'

    def __init__(self) -> None:
        self.loader = AnalyticsDataLoader()

        self.workbook = Workbook()

        default_sheet = self.workbook.active

        if default_sheet is not None:
            self.workbook.remove(default_sheet)

        self._table_counter = 0

    # ============================================================
    # MAIN BUILD
    # ============================================================

    def build(
        self,
        output_path: str | Path | None = None,
    ) -> Path:
        """
        Build complete Nexus360 Excel workbook.
        """

        self.OUTPUT_DIR.mkdir(
            parents=True,
            exist_ok=True,
        )

        if output_path is None:
            output_path = (
                self.OUTPUT_DIR
                / "Nexus360_Executive_Analytics.xlsx"
            )

        output_path = Path(output_path)

        print("[EXCEL] Loading analytics datasets...")

        datasets = self._load_datasets()

        print("[EXCEL] Building README...")
        self._build_readme()

        print("[EXCEL] Building executive dashboard...")
        self._build_executive(datasets)

        print("[EXCEL] Building monthly revenue...")
        self._build_monthly_revenue(
            datasets["monthly_revenue"]
        )

        print("[EXCEL] Building regional analytics...")
        self._build_regional(
            datasets["regional"]
        )

        print("[EXCEL] Building product analytics...")
        self._build_products(
            datasets["products"]
        )

        print("[EXCEL] Building customer 360...")
        self._build_customers(
            datasets["customers"]
        )

        print("[EXCEL] Building customer risk...")
        self._build_customer_risk(
            datasets["customer_risk"]
        )

        print("[EXCEL] Building cloud FinOps...")
        self._build_cloud_finops(
            datasets["cloud_finops"]
        )

        print("[EXCEL] Building support analytics...")
        self._build_support(
            datasets["support"]
        )

        print("[EXCEL] Building sustainability...")
        self._build_sustainability(
            datasets["sustainability"]
        )

        print("[EXCEL] Building forecast...")
        self._build_forecast()

        print("[EXCEL] Building ML predictions...")
        self._build_ml_predictions()

        print("[EXCEL] Building data dictionary...")
        self._build_dictionary(datasets)

        self._set_sheet_properties()

        self.workbook.save(output_path)

        print(
            f"[EXCEL] Workbook saved: {output_path}"
        )

        return output_path

    # ============================================================
    # DATA LOADING
    # ============================================================

    def _load_datasets(
        self,
    ) -> dict[str, pd.DataFrame]:
        """
        Load only analytics views that actually exist
        in the Nexus360 PostgreSQL semantic layer.
        """

        view_map = {
            "monthly_revenue":
                "v_monthly_revenue",

            "regional":
                "v_regional_performance",

            # Actual database view name
            "products":
                "v_product_revenue_rank",

            "customers":
                "v_customer_360",

            "customer_risk":
                "v_revenue_at_risk",

            "cloud_finops":
                "v_cloud_finops",

            "support":
                "v_support_sla_performance",

            "sustainability":
                "v_cloud_sustainability",

            "executive":
                "v_executive_kpis",
        }

        datasets: dict[str, pd.DataFrame] = {}

        for dataset_name, view_name in view_map.items():

            print(
                f"[LOAD] analytics.{view_name}"
            )

            dataframe = self.loader.load_view(
                view_name
            )

            datasets[dataset_name] = dataframe

            print(
                f"[OK]   "
                f"{dataset_name:<20} "
                f"{len(dataframe):,} rows"
            )

        return datasets

    # ============================================================
    # COMMON TITLE
    # ============================================================

    def _title(
        self,
        ws,
        title: str,
        subtitle: str,
        end_column: int = 10,
    ) -> None:

        end_column = max(
            2,
            end_column,
        )

        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=end_column,
        )

        title_cell = ws.cell(
            row=1,
            column=1,
        )

        title_cell.value = title

        title_cell.font = Font(
            size=20,
            bold=True,
            color=self.WHITE,
        )

        title_cell.fill = PatternFill(
            fill_type="solid",
            fgColor=self.NAVY,
        )

        title_cell.alignment = Alignment(
            vertical="center",
        )

        ws.row_dimensions[1].height = 32

        ws.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=end_column,
        )

        subtitle_cell = ws.cell(
            row=2,
            column=1,
        )

        subtitle_cell.value = subtitle

        subtitle_cell.font = Font(
            size=10,
            italic=True,
            color=self.MID_GRAY,
        )

        subtitle_cell.alignment = Alignment(
            vertical="center",
        )

        ws.row_dimensions[2].height = 22

    # ============================================================
    # SECTION HEADER
    # ============================================================

    def _section_header(
        self,
        ws,
        row: int,
        text: str,
        end_column: int = 10,
    ) -> None:

        end_column = max(
            2,
            end_column,
        )

        ws.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=end_column,
        )

        cell = ws.cell(
            row=row,
            column=1,
        )

        cell.value = text

        cell.font = Font(
            bold=True,
            color=self.WHITE,
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=self.BLUE,
        )

        cell.alignment = Alignment(
            vertical="center",
        )

    # ============================================================
    # SAFE EXCEL VALUE
    # ============================================================

    def _excel_safe(
        self,
        value,
    ):

        if value is None:
            return None

        if isinstance(
            value,
            pd.Timestamp,
        ):
            return value.to_pydatetime()

        if isinstance(
            value,
            np.datetime64,
        ):
            return (
                pd.Timestamp(value)
                .to_pydatetime()
            )

        if isinstance(
            value,
            np.integer,
        ):
            return int(value)

        if isinstance(
            value,
            np.floating,
        ):
            value = float(value)

            if not np.isfinite(value):
                return None

            return value

        if isinstance(
            value,
            np.bool_,
        ):
            return bool(value)

        if isinstance(
            value,
            float,
        ):
            if not np.isfinite(value):
                return None

        try:
            if pd.isna(value):
                return None
        except (TypeError, ValueError):
            pass

        return value

    # ============================================================
    # WRITE DATAFRAME
    # ============================================================

    def _write_dataframe(
        self,
        ws,
        df: pd.DataFrame,
        start_row: int = 4,
        start_col: int = 1,
        table_name: str | None = None,
    ) -> tuple[int, int]:

        data = df.copy()

        if len(data.columns) == 0:
            ws.cell(
                start_row,
                start_col,
                "No columns available",
            )

            return start_row, start_col

        # Convert datetime columns to Python date/datetime
        for column in data.columns:

            if pd.api.types.is_datetime64_any_dtype(
                data[column]
            ):
                data[column] = (
                    data[column]
                    .dt.to_pydatetime()
                )

        header_row = start_row

        # Header
        for col_index, column in enumerate(
            data.columns,
            start=start_col,
        ):

            cell = ws.cell(
                row=header_row,
                column=col_index,
            )

            cell.value = str(column)

            cell.font = Font(
                bold=True,
                color=self.WHITE,
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.BLUE,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        # Data
        for row_offset, row_values in enumerate(
            data.itertuples(
                index=False,
                name=None,
            ),
            start=1,
        ):

            excel_row = (
                header_row
                + row_offset
            )

            for col_offset, value in enumerate(
                row_values,
                start=0,
            ):

                safe_value = self._excel_safe(
                    value
                )

                ws.cell(
                    row=excel_row,
                    column=(
                        start_col
                        + col_offset
                    ),
                    value=safe_value,
                )

        end_row = (
            header_row
            + len(data)
        )

        end_col = (
            start_col
            + len(data.columns)
            - 1
        )

        # Excel Table only if rows exist
        if len(data) > 0:

            if table_name is None:

                self._table_counter += 1

                table_name = (
                    f"NexusTable"
                    f"{self._table_counter}"
                )

            table_ref = (
                f"{get_column_letter(start_col)}"
                f"{header_row}:"
                f"{get_column_letter(end_col)}"
                f"{end_row}"
            )

            table = Table(
                displayName=table_name,
                ref=table_ref,
            )

            table.tableStyleInfo = (
                TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
            )

            ws.add_table(table)

        ws.freeze_panes = (
            f"A{header_row + 1}"
        )

        self._format_columns(
            ws=ws,
            columns=data.columns,
            start_col=start_col,
            data_start_row=header_row + 1,
            end_row=end_row,
        )

        self._auto_width(ws)

        return end_row, end_col

    # ============================================================
    # FORMAT COLUMNS
    # ============================================================

    def _format_columns(
        self,
        ws,
        columns: Iterable[str],
        start_col: int,
        data_start_row: int,
        end_row: int,
    ) -> None:

        if end_row < data_start_row:
            return

        for offset, column in enumerate(
            columns
        ):

            excel_col = (
                start_col
                + offset
            )

            name = str(
                column
            ).lower()

            number_format = None

            # Percentages first because some percentage
            # fields can also contain words like cost/rate.
            if (
                "pct" in name
                or "percentage" in name
                or name.endswith("_rate")
                or "margin_pct" in name
            ):
                # Semantic views use percentage points
                # such as 95.5 rather than 0.955.
                number_format = "0.00"

            elif (
                "usd" in name
                or "revenue" in name
                or "cost" in name
                or "profit" in name
                or "contract_value" in name
            ):
                number_format = (
                    self.CURRENCY_FORMAT
                )

            elif (
                "date" in name
                or name.endswith("_month")
            ):
                number_format = (
                    self.DATE_FORMAT
                )

            elif (
                "count" in name
                or "transactions" in name
                or "customers" in name
                or "tickets" in name
                or "seats" in name
                or name.endswith("_rank")
            ):
                number_format = (
                    self.INTEGER_FORMAT
                )

            if number_format is None:
                continue

            for row_number in range(
                data_start_row,
                end_row + 1,
            ):

                ws.cell(
                    row=row_number,
                    column=excel_col,
                ).number_format = (
                    number_format
                )

    # ============================================================
    # AUTO WIDTH
    # ============================================================

    def _auto_width(
        self,
        ws,
    ) -> None:

        for column_cells in ws.columns:

            first_cell = column_cells[0]

            column_letter = (
                get_column_letter(
                    first_cell.column
                )
            )

            max_length = 0

            for cell in column_cells:

                value = cell.value

                if value is None:
                    continue

                try:
                    length = len(
                        str(value)
                    )

                    max_length = max(
                        max_length,
                        length,
                    )

                except Exception:
                    continue

            ws.column_dimensions[
                column_letter
            ].width = min(
                max(
                    max_length + 2,
                    11,
                ),
                35,
            )

    # ============================================================
    # README
    # ============================================================

    def _build_readme(
        self,
    ) -> None:

        ws = self.workbook.create_sheet(
            "00_README"
        )

        self._title(
            ws,
            "NEXUS 360",
            (
                "Enterprise Analytics & "
                "Predictive Intelligence Workbook"
            ),
            8,
        )

        generated_time = datetime.now(
            timezone.utc
        ).strftime(
            "%Y-%m-%d %H:%M:%S UTC"
        )

        content = [
            (
                "Purpose",
                (
                    "Executive analytical workbook "
                    "generated automatically from the "
                    "Nexus360 PostgreSQL semantic layer "
                    "and machine-learning outputs."
                ),
            ),
            (
                "Architecture",
                (
                    "External APIs / enterprise-style "
                    "datasets → PostgreSQL → Warehouse → "
                    "Analytics → Python/ML → Excel BI."
                ),
            ),
            (
                "Generated",
                generated_time,
            ),
            (
                "Excel Refresh",
                (
                    "python -m "
                    "scripts.build_excel_report"
                ),
            ),
            (
                "ML Refresh",
                (
                    "python -m "
                    "scripts.train_models"
                ),
            ),
            (
                "Data Scope",
                (
                    "Portfolio demonstration system "
                    "using project datasets and public "
                    "external sources. It must not be "
                    "represented as Microsoft's internal "
                    "customer or financial data."
                ),
            ),
        ]

        row = 5

        for key, value in content:

            key_cell = ws.cell(
                row=row,
                column=1,
            )

            key_cell.value = key

            key_cell.font = Font(
                bold=True,
                color=self.BLUE,
            )

            ws.merge_cells(
                start_row=row,
                start_column=2,
                end_row=row,
                end_column=8,
            )

            value_cell = ws.cell(
                row=row,
                column=2,
            )

            value_cell.value = value

            value_cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

            row += 2

        self._section_header(
            ws,
            row,
            "Workbook Navigation",
            8,
        )

        row += 1

        navigation = [
            "01_EXECUTIVE",
            "02_MONTHLY_REVENUE",
            "03_REGIONAL",
            "04_PRODUCTS",
            "05_CUSTOMERS",
            "06_CUSTOMER_RISK",
            "07_CLOUD_FINOPS",
            "08_SUPPORT",
            "09_SUSTAINABILITY",
            "10_FORECAST",
            "11_ML_PREDICTIONS",
            "12_DATA_DICTIONARY",
        ]

        for sheet_name in navigation:

            cell = ws.cell(
                row=row,
                column=1,
            )

            cell.value = sheet_name

            cell.hyperlink = (
                f"#'{sheet_name}'!A1"
            )

            cell.style = "Hyperlink"

            row += 1

        ws.column_dimensions[
            "A"
        ].width = 30

        ws.column_dimensions[
            "B"
        ].width = 35

    # ============================================================
    # EXECUTIVE
    # ============================================================

    def _build_executive(
        self,
        datasets: dict[str, pd.DataFrame],
    ) -> None:

        ws = self.workbook.create_sheet(
            "01_EXECUTIVE"
        )

        self._title(
            ws,
            (
                "NEXUS 360 — "
                "EXECUTIVE COMMAND CENTER"
            ),
            (
                "Enterprise revenue, customer, "
                "cloud and risk intelligence"
            ),
            12,
        )

        monthly = (
            datasets[
                "monthly_revenue"
            ].copy()
        )

        customers = (
            datasets[
                "customers"
            ].copy()
        )

        risk = (
            datasets[
                "customer_risk"
            ].copy()
        )

        cloud = (
            datasets[
                "cloud_finops"
            ].copy()
        )

        # --------------------------------------------------------
        # KPIs
        # --------------------------------------------------------

        total_revenue = (
            float(
                monthly[
                    "net_revenue_usd"
                ]
                .fillna(0)
                .sum()
            )
            if (
                "net_revenue_usd"
                in monthly.columns
            )
            else 0.0
        )

        total_customers = (
            int(
                customers[
                    "customer_id"
                ].nunique()
            )
            if (
                "customer_id"
                in customers.columns
            )
            else len(customers)
        )

        total_risk = (
            float(
                risk[
                    "revenue_at_risk_usd"
                ]
                .fillna(0)
                .sum()
            )
            if (
                "revenue_at_risk_usd"
                in risk.columns
            )
            else 0.0
        )

        cloud_cost = (
            float(
                cloud[
                    "estimated_cost_usd"
                ]
                .fillna(0)
                .sum()
            )
            if (
                "estimated_cost_usd"
                in cloud.columns
            )
            else 0.0
        )

        kpis = [
            (
                "Total Revenue",
                total_revenue,
                self.CURRENCY_FORMAT,
            ),
            (
                "Customers",
                total_customers,
                self.INTEGER_FORMAT,
            ),
            (
                "Revenue at Risk",
                total_risk,
                self.CURRENCY_FORMAT,
            ),
            (
                "Cloud Cost",
                cloud_cost,
                self.CURRENCY_FORMAT,
            ),
        ]

        start_columns = [
            1,
            4,
            7,
            10,
        ]

        for (
            title,
            value,
            number_format,
        ), col in zip(
            kpis,
            start_columns,
        ):

            ws.merge_cells(
                start_row=4,
                start_column=col,
                end_row=4,
                end_column=col + 1,
            )

            ws.merge_cells(
                start_row=5,
                start_column=col,
                end_row=6,
                end_column=col + 1,
            )

            title_cell = ws.cell(
                row=4,
                column=col,
            )

            title_cell.value = title

            title_cell.font = Font(
                bold=True,
                color=self.WHITE,
            )

            title_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.BLUE,
            )

            title_cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            value_cell = ws.cell(
                row=5,
                column=col,
            )

            value_cell.value = value

            value_cell.number_format = (
                number_format
            )

            value_cell.font = Font(
                size=16,
                bold=True,
                color=self.DARK_TEXT,
            )

            value_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.LIGHT_BLUE,
            )

            value_cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        # --------------------------------------------------------
        # Excel formulas
        # --------------------------------------------------------

        ws["A8"] = (
            "Risk / Revenue Ratio"
        )

        ws["A8"].font = Font(
            bold=True
        )

        ws["B8"] = (
            "=IFERROR(G5/A5,0)"
        )

        ws[
            "B8"
        ].number_format = (
            self.PERCENT_FORMAT
        )

        ws["D8"] = (
            "Revenue / Customer"
        )

        ws["D8"].font = Font(
            bold=True
        )

        ws["E8"] = (
            "=IFERROR(A5/D5,0)"
        )

        ws[
            "E8"
        ].number_format = (
            self.CURRENCY_FORMAT
        )

        # --------------------------------------------------------
        # Revenue Trend
        # --------------------------------------------------------

        required_columns = {
            "revenue_month",
            "net_revenue_usd",
        }

        if required_columns.issubset(
            monthly.columns
        ):

            trend = monthly[
                [
                    "revenue_month",
                    "net_revenue_usd",
                ]
            ].copy()

            trend[
                "revenue_month"
            ] = pd.to_datetime(
                trend[
                    "revenue_month"
                ],
                errors="coerce",
            )

            trend = (
                trend
                .dropna(
                    subset=[
                        "revenue_month"
                    ]
                )
                .sort_values(
                    "revenue_month"
                )
            )

        else:
            trend = pd.DataFrame(
                columns=[
                    "revenue_month",
                    "net_revenue_usd",
                ]
            )

        start_row = 11

        self._section_header(
            ws,
            start_row,
            "Revenue Trend",
            6,
        )

        table_header_row = (
            start_row + 1
        )

        for index, column in enumerate(
            trend.columns,
            start=1,
        ):

            cell = ws.cell(
                row=table_header_row,
                column=index,
            )

            cell.value = column

            cell.font = Font(
                bold=True,
                color=self.WHITE,
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.BLUE,
            )

        for row_number, values in enumerate(
            trend.itertuples(
                index=False,
                name=None,
            ),
            start=table_header_row + 1,
        ):

            date_value = self._excel_safe(
                values[0]
            )

            revenue_value = self._excel_safe(
                values[1]
            )

            ws.cell(
                row=row_number,
                column=1,
                value=date_value,
            ).number_format = (
                self.DATE_FORMAT
            )

            ws.cell(
                row=row_number,
                column=2,
                value=revenue_value,
            ).number_format = (
                self.CURRENCY_FORMAT
            )

        if len(trend) > 0:

            chart = LineChart()

            chart.title = (
                "Monthly Net Revenue"
            )

            chart.y_axis.title = (
                "Revenue USD"
            )

            chart.x_axis.title = (
                "Month"
            )

            chart.height = 8
            chart.width = 15

            data_reference = Reference(
                ws,
                min_col=2,
                min_row=table_header_row,
                max_row=(
                    table_header_row
                    + len(trend)
                ),
            )

            category_reference = Reference(
                ws,
                min_col=1,
                min_row=(
                    table_header_row
                    + 1
                ),
                max_row=(
                    table_header_row
                    + len(trend)
                ),
            )

            chart.add_data(
                data_reference,
                titles_from_data=True,
            )

            chart.set_categories(
                category_reference
            )

            ws.add_chart(
                chart,
                "D12",
            )

        ws.freeze_panes = "A11"

        self._auto_width(ws)

    # ============================================================
    # STANDARD DATA SHEET
    # ============================================================

    def _standard_data_sheet(
        self,
        sheet_name: str,
        title: str,
        subtitle: str,
        df: pd.DataFrame,
        table_name: str,
    ):

        ws = self.workbook.create_sheet(
            sheet_name
        )

        self._title(
            ws,
            title,
            subtitle,
            max(
                len(df.columns),
                8,
            ),
        )

        self._write_dataframe(
            ws,
            df,
            start_row=4,
            table_name=table_name,
        )

        return ws

    # ============================================================
    # MONTHLY REVENUE
    # ============================================================

    def _build_monthly_revenue(
        self,
        df: pd.DataFrame,
    ) -> None:

        ws = self._standard_data_sheet(
            "02_MONTHLY_REVENUE",
            "Monthly Revenue Analytics",
            (
                "Revenue, profitability, growth "
                "and customer economics"
            ),
            df,
            "MonthlyRevenueTable",
        )

        if (
            "net_revenue_usd"
            in df.columns
            and len(df) > 0
        ):

            column_index = (
                list(df.columns).index(
                    "net_revenue_usd"
                )
                + 1
            )

            column_letter = (
                get_column_letter(
                    column_index
                )
            )

            ws.conditional_formatting.add(
                (
                    f"{column_letter}5:"
                    f"{column_letter}"
                    f"{4 + len(df)}"
                ),
                DataBarRule(
                    start_type="min",
                    end_type="max",
                    color=self.BLUE,
                ),
            )

    # ============================================================
    # REGIONAL
    # ============================================================

    def _build_regional(
        self,
        df: pd.DataFrame,
    ) -> None:

        self._standard_data_sheet(
            "03_REGIONAL",
            "Regional Performance",
            (
                "Geographic revenue and "
                "commercial performance"
            ),
            df,
            "RegionalPerformanceTable",
        )

    # ============================================================
    # PRODUCTS
    # ============================================================

    def _build_products(
        self,
        df: pd.DataFrame,
    ) -> None:

        ws = self._standard_data_sheet(
            "04_PRODUCTS",
            "Product Portfolio Intelligence",
            (
                "Product ranking, contribution "
                "and portfolio economics"
            ),
            df,
            "ProductRevenueRankTable",
        )

        if len(df) == 0:
            return

        revenue_candidates = [
            column
            for column in df.columns
            if (
                "revenue" in column.lower()
                and "share"
                not in column.lower()
                and "rank"
                not in column.lower()
            )
        ]

        product_candidates = [
            "product_name",
            "product_code",
            "product_family",
        ]

        product_column = None

        for candidate in product_candidates:

            if candidate in df.columns:
                product_column = candidate
                break

        if (
            not revenue_candidates
            or product_column is None
        ):
            return

        revenue_column = (
            revenue_candidates[0]
        )

        revenue_col_index = (
            list(df.columns).index(
                revenue_column
            )
            + 1
        )

        product_col_index = (
            list(df.columns).index(
                product_column
            )
            + 1
        )

        max_data_rows = min(
            len(df),
            10,
        )

        chart = BarChart()

        chart.title = (
            "Top Product Performance"
        )

        chart.y_axis.title = (
            "Revenue"
        )

        chart.x_axis.title = (
            "Product"
        )

        chart.height = 8
        chart.width = 14

        data_reference = Reference(
            ws,
            min_col=revenue_col_index,
            min_row=4,
            max_row=(
                4 + max_data_rows
            ),
        )

        category_reference = Reference(
            ws,
            min_col=product_col_index,
            min_row=5,
            max_row=(
                4 + max_data_rows
            ),
        )

        chart.add_data(
            data_reference,
            titles_from_data=True,
        )

        chart.set_categories(
            category_reference
        )

        chart_column = get_column_letter(
            len(df.columns) + 2
        )

        ws.add_chart(
            chart,
            f"{chart_column}4",
        )

    # ============================================================
    # CUSTOMERS
    # ============================================================

    def _build_customers(
        self,
        df: pd.DataFrame,
    ) -> None:

        self._standard_data_sheet(
            "05_CUSTOMERS",
            "Customer 360",
            (
                "Unified commercial, support "
                "and cloud customer profile"
            ),
            df,
            "Customer360Table",
        )

    # ============================================================
    # CUSTOMER RISK
    # ============================================================

    def _build_customer_risk(
        self,
        df: pd.DataFrame,
    ) -> None:

        ws = self._standard_data_sheet(
            "06_CUSTOMER_RISK",
            "Customer Revenue Risk",
            (
                "Customer-level revenue exposure "
                "and risk indicators"
            ),
            df,
            "CustomerRiskTable",
        )

        if (
            "revenue_at_risk_usd"
            in df.columns
            and len(df) > 0
        ):

            column_index = (
                list(df.columns).index(
                    "revenue_at_risk_usd"
                )
                + 1
            )

            column_letter = (
                get_column_letter(
                    column_index
                )
            )

            ws.conditional_formatting.add(
                (
                    f"{column_letter}5:"
                    f"{column_letter}"
                    f"{4 + len(df)}"
                ),
                ColorScaleRule(
                    start_type="min",
                    start_color=(
                        self.LIGHT_GREEN
                    ),
                    mid_type="percentile",
                    mid_value=50,
                    mid_color=(
                        self.LIGHT_AMBER
                    ),
                    end_type="max",
                    end_color=(
                        self.LIGHT_RED
                    ),
                ),
            )

    # ============================================================
    # CLOUD FINOPS
    # ============================================================

    def _build_cloud_finops(
        self,
        df: pd.DataFrame,
    ) -> None:

        self._standard_data_sheet(
            "07_CLOUD_FINOPS",
            "Cloud FinOps",
            (
                "Cloud consumption, cost "
                "efficiency and reliability"
            ),
            df,
            "CloudFinOpsTable",
        )

    # ============================================================
    # SUPPORT
    # ============================================================

    def _build_support(
        self,
        df: pd.DataFrame,
    ) -> None:

        self._standard_data_sheet(
            "08_SUPPORT",
            "Support & SLA Performance",
            (
                "Service quality, resolution "
                "and SLA intelligence"
            ),
            df,
            "SupportSLATable",
        )

    # ============================================================
    # SUSTAINABILITY
    # ============================================================

    def _build_sustainability(
        self,
        df: pd.DataFrame,
    ) -> None:

        self._standard_data_sheet(
            "09_SUSTAINABILITY",
            "Cloud Sustainability",
            (
                "Carbon intensity and "
                "renewable-energy analytics"
            ),
            df,
            "SustainabilityTable",
        )

    # ============================================================
    # FORECAST
    # ============================================================

    def _build_forecast(
        self,
    ) -> None:

        path = (
            self.PREDICTION_DIR
            / "monthly_revenue_forecast.csv"
        )

        if path.exists():

            df = pd.read_csv(
                path
            )

        else:

            df = pd.DataFrame(
                columns=[
                    "forecast_month",
                    "forecast_revenue_usd",
                ]
            )

        ws = self._standard_data_sheet(
            "10_FORECAST",
            "Revenue Forecast",
            (
                "Forward-looking monthly revenue "
                "projection generated by the "
                "predictive analytics layer"
            ),
            df,
            "RevenueForecastTable",
        )

        if len(df) == 0:
            return

        forecast_column = None
        month_column = None

        for column in df.columns:

            lower = column.lower()

            if (
                forecast_column is None
                and "forecast" in lower
                and (
                    "revenue" in lower
                    or "prediction" in lower
                )
            ):
                forecast_column = column

            if (
                month_column is None
                and (
                    "month" in lower
                    or "date" in lower
                )
            ):
                month_column = column

        if forecast_column is None:

            numeric_columns = (
                df.select_dtypes(
                    include=[
                        np.number
                    ]
                ).columns.tolist()
            )

            if numeric_columns:
                forecast_column = (
                    numeric_columns[0]
                )

        if month_column is None:

            non_numeric = [
                column
                for column in df.columns
                if column
                != forecast_column
            ]

            if non_numeric:
                month_column = (
                    non_numeric[0]
                )

        if (
            forecast_column is None
            or month_column is None
        ):
            return

        forecast_col_index = (
            list(df.columns).index(
                forecast_column
            )
            + 1
        )

        month_col_index = (
            list(df.columns).index(
                month_column
            )
            + 1
        )

        chart = LineChart()

        chart.title = (
            "Revenue Forecast"
        )

        chart.y_axis.title = (
            "Forecast Revenue"
        )

        chart.x_axis.title = (
            "Forecast Period"
        )

        chart.height = 8
        chart.width = 15

        data_reference = Reference(
            ws,
            min_col=forecast_col_index,
            min_row=4,
            max_row=(
                4 + len(df)
            ),
        )

        category_reference = Reference(
            ws,
            min_col=month_col_index,
            min_row=5,
            max_row=(
                4 + len(df)
            ),
        )

        chart.add_data(
            data_reference,
            titles_from_data=True,
        )

        chart.set_categories(
            category_reference
        )

        chart_column = get_column_letter(
            len(df.columns) + 2
        )

        ws.add_chart(
            chart,
            f"{chart_column}4",
        )

    # ============================================================
    # ML PREDICTIONS
    # ============================================================

    def _build_ml_predictions(
        self,
    ) -> None:

        ws = self.workbook.create_sheet(
            "11_ML_PREDICTIONS"
        )

        self._title(
            ws,
            "Machine Learning Predictions",
            (
                "Renewal risk and revenue-at-risk "
                "predictive outputs"
            ),
            12,
        )

        renewal_path = (
            self.PREDICTION_DIR
            / "renewal_risk_predictions.csv"
        )

        revenue_risk_path = (
            self.PREDICTION_DIR
            / "revenue_at_risk_predictions.csv"
        )

        current_end_row = 5

        # --------------------------------------------------------
        # Renewal risk
        # --------------------------------------------------------

        if renewal_path.exists():

            renewal = pd.read_csv(
                renewal_path
            )

            self._section_header(
                ws,
                4,
                "Renewal Risk Predictions",
                max(
                    len(
                        renewal.columns
                    ),
                    8,
                ),
            )

            current_end_row, _ = (
                self._write_dataframe(
                    ws,
                    renewal,
                    start_row=5,
                    table_name=(
                        "RenewalPredictionTable"
                    ),
                )
            )

        else:

            ws.cell(
                row=5,
                column=1,
                value=(
                    "Renewal-risk prediction "
                    "file not found. Run "
                    "python -m scripts.train_models."
                ),
            )

        # --------------------------------------------------------
        # Revenue at risk
        # --------------------------------------------------------

        risk_section_row = (
            current_end_row + 3
        )

        if revenue_risk_path.exists():

            revenue_risk = pd.read_csv(
                revenue_risk_path
            )

            self._section_header(
                ws,
                risk_section_row,
                "Revenue-at-Risk Predictions",
                max(
                    len(
                        revenue_risk.columns
                    ),
                    8,
                ),
            )

            self._write_dataframe(
                ws,
                revenue_risk,
                start_row=(
                    risk_section_row
                    + 1
                ),
                table_name=(
                    "RevenueRiskPredictionTable"
                ),
            )

        else:

            ws.cell(
                row=risk_section_row,
                column=1,
                value=(
                    "Revenue-at-risk prediction "
                    "file not found. Run "
                    "python -m scripts.train_models."
                ),
            )

        self._auto_width(ws)

    # ============================================================
    # DATA DICTIONARY
    # ============================================================

    def _build_dictionary(
        self,
        datasets: dict[str, pd.DataFrame],
    ) -> None:

        rows = []

        for (
            dataset_name,
            dataframe,
        ) in datasets.items():

            for column in dataframe.columns:

                series = dataframe[
                    column
                ]

                rows.append(
                    {
                        "dataset":
                            dataset_name,

                        "column":
                            column,

                        "pandas_dtype":
                            str(
                                series.dtype
                            ),

                        "row_count":
                            int(
                                len(series)
                            ),

                        "non_null_rows":
                            int(
                                series
                                .notna()
                                .sum()
                            ),

                        "null_rows":
                            int(
                                series
                                .isna()
                                .sum()
                            ),

                        "unique_values":
                            int(
                                series.nunique(
                                    dropna=True
                                )
                            ),
                    }
                )

        dictionary = pd.DataFrame(
            rows
        )

        self._standard_data_sheet(
            "12_DATA_DICTIONARY",
            "Analytics Data Dictionary",
            (
                "Field-level metadata generated "
                "from Nexus360 semantic datasets"
            ),
            dictionary,
            "DataDictionaryTable",
        )

    # ============================================================
    # WORKBOOK PROPERTIES
    # ============================================================

    def _set_sheet_properties(
        self,
    ) -> None:

        for ws in self.workbook.worksheets:

            ws.sheet_view.showGridLines = (
                False
            )

            if ws.sheet_properties.pageSetUpPr is not None:
                ws.sheet_properties.pageSetUpPr.fitToPage = (
                    True
                )

            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0

            if ws.oddFooter is not None:
                ws.oddFooter.center.text = (
                    "Nexus 360 Enterprise Analytics"
                )

                ws.oddFooter.right.text = (
                    "Page &P of &N"
                )

        if self.workbook.worksheets:
            self.workbook.active = 0