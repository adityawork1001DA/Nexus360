from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.reporting.excel_report import NexusExcelReport


OUTPUT = Path(
    "reports/excel/"
    "Nexus360_Executive_Analytics_TEST.xlsx"
)


@pytest.fixture(scope="module")
def workbook_path():

    report = NexusExcelReport()

    path = report.build(
        OUTPUT
    )

    yield path

    if path.exists():
        path.unlink()


@pytest.fixture(scope="module")
def workbook(workbook_path):

    wb = load_workbook(
        workbook_path,
        data_only=False,
    )

    yield wb

    wb.close()


def test_excel_report_created(
    workbook_path,
):
    assert workbook_path.exists()


def test_excel_report_non_empty(
    workbook_path,
):
    assert workbook_path.stat().st_size > 0


def test_expected_sheet_count(
    workbook,
):
    assert len(
        workbook.sheetnames
    ) == 13


def test_expected_sheets_exist(
    workbook,
):
    expected = {
        "00_README",
        "01_EXECUTIVE",
        "02_MONTHLY_REVENUE",
        "03_REGIONAL",
        "04_PRODUCTS",
        "05_CUSTOMERS",
        "06_CUSTOMER_RISK",
        "07_CLOUD_FINOPS",
        "08_SUPPORT",
        "09_SUSTAINABILITY",
        "10_FORECAST",
        "11_ML_PREDICTIONS",
        "12_DATA_DICTIONARY",
    }

    assert expected == set(
        workbook.sheetnames
    )


def test_readme_has_title(
    workbook,
):
    ws = workbook[
        "00_README"
    ]

    assert ws["A1"].value == (
        "NEXUS 360"
    )


def test_executive_sheet_has_title(
    workbook,
):
    ws = workbook[
        "01_EXECUTIVE"
    ]

    assert "EXECUTIVE" in str(
        ws["A1"].value
    ).upper()


def test_executive_revenue_positive(
    workbook,
):
    ws = workbook[
        "01_EXECUTIVE"
    ]

    assert float(
        ws["A5"].value
    ) > 0


def test_executive_customer_count_positive(
    workbook,
):
    ws = workbook[
        "01_EXECUTIVE"
    ]

    assert int(
        ws["D5"].value
    ) > 0


def test_executive_contains_formula(
    workbook,
):
    ws = workbook[
        "01_EXECUTIVE"
    ]

    assert isinstance(
        ws["B8"].value,
        str,
    )

    assert ws[
        "B8"
    ].value.startswith("=")


def test_monthly_revenue_has_table(
    workbook,
):
    ws = workbook[
        "02_MONTHLY_REVENUE"
    ]

    assert len(
        ws.tables
    ) >= 1


def test_customer_sheet_has_table(
    workbook,
):
    ws = workbook[
        "05_CUSTOMERS"
    ]

    assert len(
        ws.tables
    ) >= 1


def test_risk_sheet_has_table(
    workbook,
):
    ws = workbook[
        "06_CUSTOMER_RISK"
    ]

    assert len(
        ws.tables
    ) >= 1


def test_cloud_finops_has_table(
    workbook,
):
    ws = workbook[
        "07_CLOUD_FINOPS"
    ]

    assert len(
        ws.tables
    ) >= 1


def test_support_has_table(
    workbook,
):
    ws = workbook[
        "08_SUPPORT"
    ]

    assert len(
        ws.tables
    ) >= 1


def test_sustainability_has_table(
    workbook,
):
    ws = workbook[
        "09_SUSTAINABILITY"
    ]

    assert len(
        ws.tables
    ) >= 1


def test_forecast_sheet_exists(
    workbook,
):
    ws = workbook[
        "10_FORECAST"
    ]

    assert ws.max_row >= 4


def test_forecast_has_chart_when_data_available(
    workbook,
):
    path = Path(
        "reports/predictions/"
        "monthly_revenue_forecast.csv"
    )

    ws = workbook[
        "10_FORECAST"
    ]

    if path.exists():
        df = pd.read_csv(path)

        if not df.empty:
            assert len(
                ws._charts
            ) >= 1


def test_ml_prediction_sheet_exists(
    workbook,
):
    ws = workbook[
        "11_ML_PREDICTIONS"
    ]

    assert ws.max_row >= 4


def test_data_dictionary_has_table(
    workbook,
):
    ws = workbook[
        "12_DATA_DICTIONARY"
    ]

    assert len(
        ws.tables
    ) == 1


def test_gridlines_disabled(
    workbook,
):
    for ws in workbook.worksheets:
        assert (
            ws.sheet_view.showGridLines
            is False
        )


def test_monthly_sheet_frozen(
    workbook,
):
    ws = workbook[
        "02_MONTHLY_REVENUE"
    ]

    assert ws.freeze_panes is not None


def test_customer_risk_conditional_formatting(
    workbook,
):
    ws = workbook[
        "06_CUSTOMER_RISK"
    ]

    assert len(
        ws.conditional_formatting
    ) >= 1


def test_workbook_navigation_links(
    workbook,
):
    ws = workbook[
        "00_README"
    ]

    hyperlinks = []

    for row in ws.iter_rows():
        for cell in row:
            if cell.hyperlink:
                hyperlinks.append(
                    cell.hyperlink.target
                )

    assert len(
        hyperlinks
    ) >= 10